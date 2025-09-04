from pptx import Presentation
import json
from typing import Dict, List, Optional
import copy


class SlideGenerator:
    def __init__(self, template_path: str, schema_path: str):
        """
        Initialize the SlideGenerator with a template presentation and schema.
        
        Args:
            template_path (str): Path to the template PPTX file
            schema_path (str): Path to the template schema JSON file
        """
        self.prs = Presentation(template_path)
        with open(schema_path, 'r') as f:
            self.schema = json.load(f)
        
        # Create template slide mapping (slides are in order in template.pptx)
        self.template_slides = {
            slide_info['name']: idx 
            for idx, slide_info in enumerate(self.schema['slides'])
        }
    
    def _get_template_slide(self, template_name: str) -> int:
        """Get the slide index for a given template name."""
        if template_name not in self.template_slides:
            raise ValueError(f"Template '{template_name}' not found in schema")
        return self.template_slides[template_name]
    
    def _duplicate_slide(self, template_name: str) -> int:
        """
        Duplicate a slide from the template and return its new index.
        
        Args:
            template_name (str): Name of the template slide to duplicate
            
        Returns:
            int: Index of the new slide
        """
        source_slide = self.prs.slides[self._get_template_slide(template_name)]
        blank_slide_layout = source_slide.slide_layout
        new_slide = self.prs.slides.add_slide(blank_slide_layout)
        
        # Copy all shapes from the source slide
        for shape in source_slide.shapes:
            el = shape.element
            new_el = copy.deepcopy(el)
            new_slide.shapes._spTree.insert_element_before(new_el, 'p:extLst')
        
        return len(self.prs.slides) - 1
    
    def _update_slide_content(self, slide_idx: int, content: Dict):
        """
        Update the content of a slide based on the provided data.
        
        Args:
            slide_idx (int): Index of the slide to update
            content (dict): Content to insert into the slide
        """
        slide = self.prs.slides[slide_idx]
        
        for shape in slide.shapes:
            if not shape.has_text_frame:
                continue
                
            text_frame = shape.text_frame
            text = text_frame.text
            
            # Check if this is a bullet point placeholder
            if "{{ bullet_points }}" in text and "bullet_points" in content:
                # Clear existing text
                text_frame.clear()
                # Add bullet points
                for point in content["bullet_points"]:
                    p = text_frame.add_paragraph()
                    p.text = point
                    p.level = 0
            else:
                # Handle regular text replacement
                for paragraph in text_frame.paragraphs:
                    for key, value in content.items():
                        if key != "bullet_points":  # Skip bullet points here
                            if f"{{{{ {key} }}}}" in paragraph.text:
                                paragraph.text = paragraph.text.replace(f"{{{{ {key} }}}}", str(value))
    
    def generate_presentation(self, config_path: str, output_path: str):
        """
        Generate a new presentation based on the configuration.
        
        Args:
            config_path (str): Path to the configuration JSON file
            output_path (str): Path where the generated presentation should be saved
        """
        with open(config_path, 'r') as f:
            config = json.load(f)
        
        # Store original template slides count
        template_count = len(self.template_slides)
        
        # Create new slides based on configuration
        for section in config['sections']:
            for slide in section['slides']:
                template_name = slide['template']
                # Duplicate template slide - it will be added at the end
                new_slide_idx = self._duplicate_slide(template_name)
                
                # Update content for the new slide if provided
                if 'content' in slide:
                    self._update_slide_content(new_slide_idx, slide['content'])
        
        # Remove original template slides
        for _ in range(template_count):
            xml_slides = self.prs.slides._sldIdLst
            if xml_slides:
                xml_slides.remove(xml_slides[0])
        
        # Save the presentation
        self.prs.save(output_path)
